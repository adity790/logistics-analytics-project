"""Small python-docx toolkit: consistent styling for the four weekly reports."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREY = RGBColor(0x52, 0x51, 0x4E)
HEADER_FILL = "1F3864"
ZEBRA_FILL = "F2F6FB"
CODE_FILL = "F4F4F2"
CALLOUT_FILL = "E8F1FC"


def _shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=50, bottom=50, left=80, right=80) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tcPr.append(mar)


def _set_borders(table, color="BFC7D5", sz=4) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _add_field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)


class Report:
    def __init__(self, short_title: str):
        self.doc = Document()
        self.short_title = short_title
        self.fig_n = 0
        self.tab_n = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        self._styles()
        self._header_footer(sec)

    # ------------------------------------------------------------------ styles -----------
    def _styles(self) -> None:
        st = self.doc.styles
        n = st["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
        for name, size, color, before, after in (("Heading 1", 17, NAVY, 18, 6), ("Heading 2", 13.5, NAVY, 12, 4), ("Heading 3", 11.5, BLUE, 8, 3)):
            h = st[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
            h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True
        for name in ("List Bullet",):
            b = st[name]; b.font.name = "Calibri"; b.font.size = Pt(10.5); b.paragraph_format.space_after = Pt(3)
        from docx.enum.style import WD_STYLE_TYPE
        code = st.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = st["Normal"]; code.font.name = "Consolas"; code.font.size = Pt(8)
        code.element.rPr.rFonts.set(qn("w:ascii"), "Consolas"); code.element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        code.element.rPr.rFonts.set(qn("w:cs"), "Consolas")
        code.paragraph_format.space_after = Pt(0); code.paragraph_format.space_before = Pt(0); code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.left_indent = Cm(0.2)
        pPr = code.element.get_or_add_pPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), CODE_FILL); pPr.append(shd)
        cap = st.add_style("Cap", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = st["Normal"]; cap.font.size = Pt(9); cap.font.italic = True; cap.font.color.rgb = GREY
        cap.paragraph_format.space_after = Pt(10); cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _header_footer(self, sec) -> None:
        hp = sec.header.paragraphs[0]
        hp.text = self.short_title
        hp.runs[0].font.size = Pt(8.5); hp.runs[0].font.color.rgb = GREY
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Page "); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        r2 = fp.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = GREY; _add_field(r2, "PAGE")

    # ------------------------------------------------------------------ text -------------
    def _runs(self, par, text: str, size=None, color=None, italic=False) -> None:
        for i, chunk in enumerate(re.split(r"\*\*", text)):
            if not chunk:
                continue
            r = par.add_run(chunk); r.bold = bool(i % 2)
            r.italic = italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color

    def title_page(self, kicker: str, title: str, subtitle: str, author: str, meta: list[tuple[str, str]]) -> None:
        for _ in range(5):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph(); self._runs(p, kicker.upper(), size=11, color=BLUE); p.runs[0].bold = True
        p = self.doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(4)
        p = self.doc.add_paragraph(); r = p.add_run(subtitle); r.font.size = Pt(14); r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(30)
        t = self.doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for k, v in [("Prepared by", author)] + meta:
            row = t.add_row().cells
            row[0].width, row[1].width = Cm(3.6), Cm(11.5)
            row[0].paragraphs[0].add_run(k).bold = True
            row[1].paragraphs[0].add_run(v)
            for c in row:
                c.paragraphs[0].paragraph_format.space_after = Pt(3)
        self.pagebreak()

    def h1(self, text): return self.doc.add_heading(text, level=1)
    def h2(self, text): return self.doc.add_heading(text, level=2)
    def h3(self, text): return self.doc.add_heading(text, level=3)

    def p(self, text: str, italic=False, align=None):
        par = self.doc.add_paragraph(); self._runs(par, text, italic=italic)
        if align == "justify": par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return par

    def bullets(self, items: list[str]) -> None:
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet"); self._runs(par, it)

    def pagebreak(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def callout(self, title: str, text: str) -> None:
        t = self.doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; _shade_cell(c, CALLOUT_FILL); _cell_margins(c, 90, 90, 140, 140)
        par = c.paragraphs[0]; r = par.add_run(title + "  "); r.bold = True; r.font.color.rgb = NAVY
        self._runs(par, text)
        par.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------ tables / figures --
    def table(self, header: list[str], rows: list[list], widths_cm: list[float] | None = None, caption: str | None = None,
              font=9, align_right_from: int | None = None) -> None:
        if caption:
            self.tab_n += 1
            cp = self.doc.add_paragraph(); cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(3)
            r = cp.add_run(f"Table {self.tab_n}. "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = NAVY
            r2 = cp.add_run(caption); r2.font.size = Pt(9); r2.font.color.rgb = GREY
        t = self.doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _set_borders(t)
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; _shade_cell(c, HEADER_FILL); _cell_margins(c)
            par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
            r = par.add_run(str(h)); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if align_right_from is not None and i >= align_right_from:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        # repeat header row on page break
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); t.rows[-1]._tr.get_or_add_trPr().append(cant)
            for i, v in enumerate(row):
                c = cells[i]; _cell_margins(c)
                if ri % 2 == 1: _shade_cell(c, ZEBRA_FILL)
                par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
                self._runs(par, str(v), size=font)
                if align_right_from is not None and i >= align_right_from:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def fig(self, path: Path, caption: str, width_in: float = 6.0) -> None:
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True; par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(str(path), width=Inches(width_in))
        self.fig_n += 1
        c = self.doc.add_paragraph(style="Cap")
        r = c.add_run(f"Figure {self.fig_n}. "); r.bold = True; r.italic = False
        c.add_run(caption)

    def code(self, code: str, caption: str | None = None) -> None:
        if caption:
            cp = self.doc.add_paragraph(); cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(2)
            r = cp.add_run(caption); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = NAVY
        lines = code.rstrip("\n").split("\n")
        for ln in lines:
            par = self.doc.add_paragraph(style="CodeBlock")
            par.add_run(ln.replace("\t", "    ") if ln else " ")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def references(self, refs: list[str]) -> None:
        self.h1("References")
        for r in refs:
            par = self.doc.add_paragraph(); par.paragraph_format.left_indent = Cm(0.6); par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(3)
            self._runs(par, r, size=9.5)

    def save(self, path: Path) -> None:
        self.doc.save(str(path))
